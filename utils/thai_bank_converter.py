# -*- coding: utf-8 -*-
"""
Thai Bank Statement Converter for Odoo
Converts Thai bank statement formats (Excel/CSV) to Odoo-compatible format
"""

import pandas as pd
import openpyxl
from datetime import datetime
import base64
import io
from odoo.exceptions import UserError


class ThaiBankStatementConverter:
    """Converter for Thai bank statements to Odoo format"""
    
    def __init__(self):
        self.dataframes = []
    
    def detect_statement_format(self, df_raw, filename):
        """
        Detect the bank statement format to use specialized parser
        Returns: format_type string or None for generic parser
        """
        # Check first few rows for format signatures
        first_10_rows = []
        for i in range(min(10, len(df_raw))):
            row_data = df_raw.iloc[i].dropna().tolist()
            if row_data:
                first_10_rows.append(' '.join([str(x) for x in row_data[:5]]))
        
        combined = ' '.join(first_10_rows).lower()
        
        # SCB Saving Account format - has account metadata section
        if 'ชื่อบัญชี' in combined and 'ประเภทบัญชี' in combined and 'ออมทรัพย์' in combined:
            return 'scb_saving'
        
        # Add more format detections here as needed
        
        return 'generic'
    
    def find_column_by_keywords(self, columns, keywords):
        """Find a column that matches any of the keywords (substring match)"""
        for col in columns:
            col_lower = str(col).lower()
            for keyword in keywords:
                if keyword in col_lower:
                    return col
        return None
    
    def parse_date(self, date_str):
        """Parse various date formats and return Python date object"""
        if pd.isna(date_str) or date_str == '':
            return None
        
        date_str = str(date_str).strip()
        
        # List of date formats to try
        formats = [
            '%d-%b-%Y',      # 05-Jan-2026
            '%d-%b-%y',      # 01-Nov-25
            '%d/%m/%Y',      # 07/10/2025
            '%d/%m/%y',      # 07/10/25
            '%d-%m-%Y',      # 01-11-2025
            '%d-%m-%y',      # 01-11-25
            '%Y-%m-%d',      # 2025-11-01
            '%Y/%m/%d',      # 2025/11/01
            '%d.%m.%Y',      # 01.11.2025
            '%d.%m.%y',      # 01.11.25
            '%d %b %Y',      # 01 Nov 2025
            '%d %B %Y',      # 01 November 2025
        ]
        
        for fmt in formats:
            try:
                parsed_date = datetime.strptime(date_str, fmt)
                return parsed_date.date()
            except ValueError:
                continue
        
        # Try to handle Thai Buddhist year (BE) - convert to AD
        try:
            if '/' in date_str or '-' in date_str:
                parts = date_str.replace('-', '/').split('/')
                if len(parts) >= 3:
                    day, month, year = parts[0], parts[1], parts[2]
                    year_int = int(year)
                    # Thai year is 543 years ahead of AD
                    if year_int > 2500:
                        year_int -= 543
                    parsed_date = datetime(int(year_int), int(month), int(day))
                    return parsed_date.date()
        except:
            pass
        
        raise UserError(f"Could not parse date '{date_str}'")
    
    def clean_float_value(self, value):
        """Clean and convert a value to float, handling various formats"""
        if pd.isna(value) or value == '':
            return 0.0
        
        value_str = str(value).strip()
        if not value_str or value_str.lower() == 'nan':
            return 0.0
        
        # Remove commas and spaces
        value_str = value_str.replace(',', '').replace(' ', '')
        
        try:
            return float(value_str)
        except:
            return 0.0
    
    def is_valid_transaction_row(self, row, date_col, debit_col, credit_col, desc_col):
        """Check if a row is a valid transaction"""
        # Must have a valid date
        date_value = row.get(date_col, '') if date_col else ''
        if pd.isna(date_value) or str(date_value).strip() == '':
            return False
        
        # Check if date can be parsed
        try:
            self.parse_date(date_value)
        except:
            return False
        
        # Must have at least one amount column
        if not debit_col and not credit_col:
            # If neither debit nor credit column specified, row is still valid
            # (amount will be determined from a generic amount column later)
            return True
        
        # Must have at least one non-zero amount (debit or credit)
        debit = self.clean_float_value(row.get(debit_col, 0)) if debit_col else 0
        credit = self.clean_float_value(row.get(credit_col, 0)) if credit_col else 0
        
        if debit == 0 and credit == 0:
            return False
        
        # Skip common summary/total rows by checking description
        if desc_col:
            desc = str(row.get(desc_col, '')).lower()
            skip_keywords = ['total', 'subtotal', 'summary', 'balance', 'closing', 'opening',
                           'รวม', 'ยอดรวม', 'ยอด', 'สรุป', 'ปิด', 'เปิด']
            for keyword in skip_keywords:
                if keyword in desc:
                    return False
        
        return True
    
    def find_all_columns_by_keywords(self, columns, keywords):
        """Find all columns that match any of the keywords (substring match)"""
        matching_cols = []
        for col in columns:
            col_lower = str(col).lower()
            for keyword in keywords:
                if keyword in col_lower and col not in matching_cols:
                    matching_cols.append(col)
                    break
        return matching_cols
    
    def convert_generic(self, df):
        """Convert any bank statement format to Odoo format"""
        # Find date column
        date_col = self.find_column_by_keywords(df.columns, [
            'date', 'transaction date', 'effective date', 'วันที่', 'วันทำ', 'เวลา'
        ])
        
        # Find ALL description/detail/cheque related columns
        desc_keywords = [
            'description', 'detail', 'transaction description', 'tr description',
            'memo', 'รายละเอียด', 'หมายเหตุ', 'รายการ', 'รายละเอียดรายการ',
            'label', 'additional', 'enrichment', 'remarks'
        ]
        desc_cols = self.find_all_columns_by_keywords(df.columns, desc_keywords)
        
        cheque_keywords = [
            'cheque', 'cheque no', 'cheque number', 'check', 'check no',
            'เช็ค', 'เลขเช็ค', 'เล่มเล้ว'
        ]
        cheque_cols = self.find_all_columns_by_keywords(df.columns, cheque_keywords)
        
        # Check for "Debit/Credit" indicator column (some banks use this format)
        debit_credit_indicator_col = self.find_column_by_keywords(df.columns, [
            'debit/credit', 'dr/cr', 'transaction type'
        ])
        
        # Find debit/withdrawal column (exclude indicators)
        debit_col = self.find_column_by_keywords(df.columns, [
            'withdrawal', 'ถอน', 'จ่าย', 'ออก'
        ])
        # Only look for 'debit' if it's not the indicator column
        if not debit_col:
            for col in df.columns:
                col_lower = str(col).lower()
                if 'debit' in col_lower and col != debit_credit_indicator_col:
                    debit_col = col
                    break
        
        # Find credit/deposit column (exclude indicators)
        credit_col = self.find_column_by_keywords(df.columns, [
            'deposit', 'ฝาก', 'รับ', 'เข้า'
        ])
        # Only look for 'credit' if it's not the indicator column
        if not credit_col:
            for col in df.columns:
                col_lower = str(col).lower()
                if 'credit' in col_lower and col != debit_credit_indicator_col:
                    credit_col = col
                    break
        
        # If neither debit nor credit found, try to find a generic amount column
        amount_col = None
        if not debit_col and not credit_col:
            amount_col = self.find_column_by_keywords(df.columns, [
                'amount', 'จำนวน'
            ])
        
        if not date_col:
            raise UserError("Could not find date column in the file")
        
        if not debit_col and not credit_col and not amount_col:
            raise UserError("Could not find amount columns in the file")
        
        result = []
        skipped_rows = 0
        
        for idx, row in df.iterrows():
            # Validate if this is a transaction row
            if not self.is_valid_transaction_row(row, date_col, debit_col or amount_col, credit_col, desc_cols[0] if desc_cols else None):
                skipped_rows += 1
                continue
            
            # Parse date to Python date object
            date_value = row.get(date_col, '') if date_col else ''
            try:
                date = self.parse_date(date_value)
            except UserError:
                skipped_rows += 1
                continue
            
            # Build comprehensive label from all relevant columns
            label_parts = []
            seen_values = set()
            
            # Add cheque numbers
            for cheque_col in cheque_cols:
                cheque = str(row.get(cheque_col, '')).strip()
                if cheque and cheque != '0' and cheque.lower() != 'nan' and cheque not in seen_values:
                    label_parts.append(f"Cheque #{cheque}")
                    seen_values.add(cheque)
            
            # Add descriptions
            for desc_col in desc_cols:
                desc = str(row.get(desc_col, '')).strip()
                if desc and desc.lower() != 'nan' and desc not in seen_values:
                    label_parts.append(desc)
                    seen_values.add(desc)
            
            description = ' / '.join(label_parts) if label_parts else 'Bank Transaction'
            
            # Get debit and credit values
            if amount_col:
                # Single amount column with optional debit/credit indicator
                amount_value = self.clean_float_value(row.get(amount_col, 0))
                
                # Check if there's a debit/credit indicator column
                if debit_credit_indicator_col:
                    indicator = str(row.get(debit_credit_indicator_col, '')).lower().strip()
                    if 'debit' in indicator or 'dr' in indicator or 'withdrawal' in indicator:
                        # Debit = money out (negative)
                        debit = abs(amount_value)
                        credit = 0.0
                    else:
                        # Credit = money in (positive)
                        credit = abs(amount_value)
                        debit = 0.0
                else:
                    # No indicator - use sign of amount
                    if amount_value >= 0:
                        credit = amount_value
                        debit = 0.0
                    else:
                        debit = abs(amount_value)
                        credit = 0.0
            else:
                debit = self.clean_float_value(row.get(debit_col, 0)) if debit_col else 0.0
                credit = self.clean_float_value(row.get(credit_col, 0)) if credit_col else 0.0
                
                # Handle negative values (convert negative to positive for appropriate column)
                if debit < 0:
                    credit = abs(debit)
                    debit = 0.0
                if credit < 0:
                    debit = abs(credit)
                    credit = 0.0
            
            # Calculate signed amount for Odoo (positive = deposit, negative = withdrawal)
            amount = credit - debit
            
            result.append({
                'date': date,
                'payment_ref': description,
                'amount': amount,
            })
        
        return result
    
    def convert_scb_saving(self, df_raw):
        """
        Convert SCB Saving Account statement format
        This format has account metadata in rows 0-6, then header at row 7
        """
        # Read with header at row 7
        try:
            df = pd.read_excel(io.BytesIO(self.current_file_data), header=7)
            # Clean column names
            df.columns = df.columns.str.strip()
        except:
            # Fallback: manually extract from row 7 onwards
            header_row = 7
            columns = df_raw.iloc[header_row].tolist()
            df = df_raw.iloc[header_row+1:].copy()
            df.columns = columns
            df.columns = [str(c).strip() for c in df.columns]
        
        # SCB Saving format columns: วันที่/เวลา, ถอน, ฝาก, ยอดเงินในบัญชี, etc.
        date_col = 'วันที่/เวลา'
        debit_col = 'ถอน'  # Withdrawal
        credit_col = 'ฝาก'  # Deposit
        desc_col = 'คำอธิบายรายละเอียด' if 'คำอธิบายรายละเอียด' in df.columns else None
        
        result = []
        for idx, row in df.iterrows():
            # Get date value
            date_value = row.get(date_col, '')
            if pd.isna(date_value) or str(date_value).strip() == '':
                continue
            
            # Skip special rows like "B/F" (brought forward)
            if 'b/f' in str(date_value).lower() or 'ยอดเงินคงเหลือยกมา' in str(row.values).lower():
                continue
            
            # Parse date (format: "02/10/2025\n09:30:49")
            try:
                date_str = str(date_value).split('\n')[0].strip()  # Take only date part
                date = self.parse_date(date_str)
            except:
                continue
            
            # Get amounts
            debit = self.clean_float_value(row.get(debit_col, 0))
            credit = self.clean_float_value(row.get(credit_col, 0))
            
            if debit == 0 and credit == 0:
                continue
            
            # Build description
            desc = str(row.get(desc_col, '')).strip() if desc_col else ''
            if desc and desc.lower() != 'nan':
                description = desc
            else:
                description = 'SCB Transaction'
            
            # Calculate net amount (credit - debit)
            amount = credit - debit
            
            result.append({
                'date': date,
                'payment_ref': description,
                'amount': amount,
            })
        
        return result
    
    def find_header_row(self, df):
        """Try to find the actual header row in a messy file"""
        best_idx = 0
        best_keyword_count = 0
        
        for idx, row in df.iterrows():
            # Skip completely empty rows
            non_null_values = [x for x in row if pd.notna(x) and str(x).strip() != '']
            if len(non_null_values) < 2:
                continue
            
            row_str = ' '.join([str(x).lower() for x in row if pd.notna(x) and x != ''])
            
            # Check if row contains common header keywords
            header_keywords = ['date', 'description', 'detail', 'debit', 'credit', 'withdrawal', 
                             'deposit', 'amount', 'cheque', 'check', 'balance', 'pay',
                             'วัน', 'รายละเอียด', 'จ่าย', 'รับ', 'เบิก', 'วิธี', 'เล่มที่', 'เล่มเล้ว', 'รายการ',
                             'ถอน', 'ฝาก', 'เวลา', 'ช่องทาง']
            
            keyword_count = sum(1 for keyword in header_keywords if keyword in row_str)
            
            # Track the row with the most keyword matches
            if keyword_count > best_keyword_count:
                best_keyword_count = keyword_count
                best_idx = idx
        
        # Only return a header row if we found at least 2 keywords
        if best_keyword_count >= 2:
            return best_idx
        
        return 0
    
    def read_excel_with_merged_headers(self, file_data, header_idx):
        """Read Excel file handling merged cells in headers"""
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_data))
            ws = wb.active
            
            # Read first header row
            header_row_1 = []
            for col in range(1, ws.max_column + 1):
                cell_value = ws.cell(header_idx + 1, col).value
                header_row_1.append(str(cell_value) if cell_value else '')
            
            # Read second header row (sub-columns)
            header_row_2 = []
            if header_idx + 2 <= ws.max_row:
                for col in range(1, ws.max_column + 1):
                    cell_value = ws.cell(header_idx + 2, col).value
                    header_row_2.append(str(cell_value) if cell_value else '')
            
            # Combine headers
            final_columns = []
            for col_idx, (h1, h2) in enumerate(zip(header_row_1, header_row_2)):
                h1 = h1.strip() if h1 else ''
                h2 = h2.strip() if h2 else ''
                
                if h1 and h2 and h1 != h2 and h2.lower() != 'unnamed':
                    combined = f"{h1} {h2}"
                elif h1:
                    combined = h1
                elif h2:
                    combined = h2
                else:
                    combined = f"Unnamed: {col_idx}"
                
                final_columns.append(combined)
            
            # Read data starting from row after headers
            data_rows = []
            for row_idx in range(header_idx + 3, ws.max_row + 1):
                row_data = []
                for col in range(1, ws.max_column + 1):
                    cell_value = ws.cell(row_idx, col).value
                    row_data.append(cell_value)
                data_rows.append(row_data)
            
            df = pd.DataFrame(data_rows, columns=final_columns)
            wb.close()
            return df
            
        except Exception as e:
            return None
    
    def convert_file(self, file_data, filename):
        """
        Convert a bank statement file to Odoo format
        
        :param file_data: Binary file data (bytes)
        :param filename: Original filename
        :return: List of transaction dictionaries ready for Odoo
        """
        # Store file data for specialized parsers
        self.current_file_data = file_data
        
        # Determine file type
        file_ext = filename.lower().split('.')[-1]
        
        try:
            if file_ext == 'csv':
                df_raw = pd.read_csv(io.BytesIO(file_data), header=None)
            elif file_ext in ['xlsx', 'xls']:
                df_raw = pd.read_excel(io.BytesIO(file_data), header=None)
            else:
                raise UserError(f"Unsupported file format: {file_ext}. Please upload CSV or Excel files.")
        except Exception as e:
            raise UserError(f"Error reading file: {str(e)}")
        
        # Detect statement format
        format_type = self.detect_statement_format(df_raw, filename)
        
        # Use specialized parser if format is detected
        if format_type == 'scb_saving':
            result = self.convert_scb_saving(df_raw)
        else:
            # Generic parser for other formats
            try:
                if file_ext == 'csv':
                    df = pd.read_csv(io.BytesIO(file_data), encoding='utf-8')
                    df.columns = df.columns.str.strip()
                elif file_ext in ['xlsx', 'xls']:
                    # Find the actual header row
                    header_idx = self.find_header_row(df_raw)
                    
                    if header_idx > 0:
                        # Try to read with merged header handling
                        df_merged = self.read_excel_with_merged_headers(file_data, header_idx)
                        if df_merged is not None:
                            df = df_merged
                        else:
                            df = pd.read_excel(io.BytesIO(file_data), header=header_idx)
                    else:
                        df = pd.read_excel(io.BytesIO(file_data))
                    
                    # Clean column names (strip whitespace)
                    df.columns = df.columns.str.strip()
            except Exception as e:
                raise UserError(f"Error reading file: {str(e)}")
            
            # Convert using generic format
            result = self.convert_generic(df)
        
        if not result:
            raise UserError(
                "No valid transactions found in the file. "
                "Please ensure the file contains transaction data with dates and amounts."
            )
        
        # Sort by date
        result = sorted(result, key=lambda x: x['date'])
        
        return result
